import numpy as np
from openpyxl import load_workbook
import pandas as pd




def compute_voi_matrix(time_matrix, n_states, a_ws):
    n_scenarios, n_nodes = time_matrix.shape

    # Step 1: 将连续时间映射为状态 s ∈ [0,6]
    detection_states = np.digitize(time_matrix, detection_bins) - 1  # 变成 0~6
    VOI_matrix = np.zeros((n_nodes, n_nodes))

    for i in range(n_nodes):  # sensor放在哪个节点
        for j in range(n_nodes):  # 被观测的节点
            V_m_total = 0
            for m in range(n_states):  # message m，对应于延迟 w
                # 计算 P(s | m): 统计在sensor i检测为 w=m 时，node j 的真实检测状态分布
                mask = (detection_states[:, i] == m)
                if np.sum(mask) == 0:
                    continue
                state_j_given_m = detection_states[mask, j]
                p_s_given_m = np.bincount(state_j_given_m, minlength=n_states) / np.sum(mask)

                # 计算 E_m
                E_m = np.sum(a_ws[m] * p_s_given_m)

                # 计算 E_n：使用 node j 的先验检测状态概率
                state_j = detection_states[:, j]
                p_s = np.bincount(state_j, minlength=n_states) / len(state_j)
                E_n = np.sum(a_ws[m] * p_s)

                V_m = E_m - E_n
                p_m = np.mean(mask)
                V_m_total += p_m * V_m

            VOI_matrix[i, j] = V_m_total

    return VOI_matrix / np.max(VOI_matrix)


def compute_te_matrix(time_matrix, n_states):
    n_scenarios, n_nodes = time_matrix.shape
    detection_states = np.digitize(time_matrix, detection_bins) - 1

    TE_matrix = np.zeros((n_nodes, n_nodes))

    for i in range(n_nodes):
        state_i = detection_states[:, i]
        p_i = np.bincount(state_i, minlength=n_states) / len(state_i)

        for j in range(n_nodes):
            state_j = detection_states[:, j]
            p_j = np.bincount(state_j, minlength=n_states) / len(state_j)

            joint = np.zeros((n_states, n_states))
            for s in range(n_states):
                for t in range(n_states):
                    joint[s, t] = np.sum((state_i == s) & (state_j == t))
            joint /= len(state_i)

            # mutual information
            te = 0
            for s in range(n_states):
                for t in range(n_states):
                    if joint[s, t] > 0:
                        te += joint[s, t] * np.log(joint[s, t] / (p_i[s] * p_j[t] + 1e-12))
            TE_matrix[i, j] = te

    return TE_matrix / np.max(TE_matrix)

a_ws = np.array([
    [0, -1728, -3024, -6220, -27734, -98410, -172368],
    [-1728, 0, -1296, -4492, -26006, -96682, -170640],
    [-3024, -1296, 0, -3196, -24710, -95386, -169344],
    [-6220, -4492, -3196, 0, -21514, -92189, -166147],
    [-27734, -26006, -24710, -21514, 0, -70675, -144633],
    [-98410, -96682, -95386, -92189, -70675, 0, -73958],
    [-172368, -170640, -169344, -166147, -144633, -73958, 0]])

# 论文中的检测时间段划分（按小时）
detection_bins = [0, 1.1, 2.1, 4.1, 12.1, 24.1, 60.1, 10000]  # 共7段 => s1 ~ s7
n_states =7
# 打开工作簿和工作表
wb = load_workbook(r"D:\A_科研相关\小论文2_合作监测点布置\覆盖双目标优化\L-town模拟验证\Khorshidi_all_detect_time.xlsx", data_only=True)
ws = wb["Sheet1"]
cell_range = ws["B2:ADC3129"]
time_matrix= np.array([[cell.value for cell in row] for row in cell_range])

VOI_matrix=compute_voi_matrix(time_matrix, n_states,a_ws)
VOI_matrix_df = pd.DataFrame(VOI_matrix)
VOI_matrix_df.to_excel(r"D:\A_科研相关\小论文2_合作监测点布置\覆盖双目标优化\L-town模拟验证\Khorshidi_VOI_matrix.xlsx", index=False, header=False)

TE_matrix=compute_te_matrix(time_matrix, n_states)
TE_matrix_df = pd.DataFrame(TE_matrix)
TE_matrix_df.to_excel(r"D:\A_科研相关\小论文2_合作监测点布置\覆盖双目标优化\L-town模拟验证\Khorshidi_TE_matrix.xlsx", index=False, header=False)
