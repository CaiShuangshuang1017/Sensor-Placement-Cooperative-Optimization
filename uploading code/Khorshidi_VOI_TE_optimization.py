import numpy as np
import pandas as pd
import random
import os
from openpyxl import load_workbook
import matplotlib.pyplot as plt

def initialize_population(pop_size, n_vars, sensor_num):
    population = []
    for _ in range(pop_size):
        individual = np.zeros(n_vars, dtype=int)
        individual[np.random.choice(n_vars, sensor_num, replace=False)] = 1
        population.append(individual)
    return np.array(population)

def evaluate_individual(individual, VOI_matrix, TE_matrix):
    selected = np.where(individual == 1)[0]
    if len(selected) == 0:
        return 0, 0

    # VOI：每列最大值求和
    voi = np.sum(np.max(VOI_matrix[selected, :], axis=0))

    # TE：两两组合互信息和
    te = 0
    for i in range(len(selected)):
        for j in range(i + 1, len(selected)):
            te += TE_matrix[selected[i], selected[j]]

    return voi, te

def dominates(a, b):
    return (a[0] >= b[0] and a[1] <= b[1]) and (a[0] > b[0] or a[1] < b[1])

def non_dominated_sort(pop_objs):
    fronts = []
    S = [[] for _ in range(len(pop_objs))]
    n = [0 for _ in range(len(pop_objs))]
    rank = [0 for _ in range(len(pop_objs))]

    front = []

    for p in range(len(pop_objs)):
        for q in range(len(pop_objs)):
            if dominates(pop_objs[p], pop_objs[q]):
                S[p].append(q)
            elif dominates(pop_objs[q], pop_objs[p]):
                n[p] += 1
        if n[p] == 0:
            rank[p] = 0
            front.append(p)

    fronts.append(front)

    i = 0
    while fronts[i]:
        next_front = []
        for p in fronts[i]:
            for q in S[p]:
                n[q] -= 1
                if n[q] == 0:
                    rank[q] = i + 1
                    next_front.append(q)
        i += 1
        fronts.append(next_front)

    return fronts[:-1]

def crossover(parent1, parent2):
    point = random.randint(1, len(parent1) - 2)
    child1 = np.concatenate((parent1[:point], parent2[point:]))
    child2 = np.concatenate((parent2[:point], parent1[point:]))
    return child1, child2

def mutate(individual, mutation_rate=0.1):
    for i in range(len(individual)):
        if random.random() < mutation_rate:
            individual[i] = 1 - individual[i]
    return individual

def enforce_fixed_sensor_count(individual, sensor_num):
    current = np.sum(individual)
    if current < sensor_num:
        zeros = np.where(individual == 0)[0]
        individual[np.random.choice(zeros, sensor_num - current, replace=False)] = 1
    elif current > sensor_num:
        ones = np.where(individual == 1)[0]
        individual[np.random.choice(ones, current - sensor_num, replace=False)] = 0
    return individual

def run_fixed_sensor_optimization(VOI_matrix, TE_matrix, sensor_num, pop_size=500, generations=300):
    n_vars = VOI_matrix.shape[0]
    population = initialize_population(pop_size, n_vars, sensor_num)

    for gen in range(generations):
        fitness = [evaluate_individual(ind, VOI_matrix, TE_matrix) for ind in population]
        fronts = non_dominated_sort(fitness)

        new_population = []
        for front in fronts:
            if len(new_population) + len(front) > pop_size:
                break
            new_population.extend([population[i] for i in front])

        while len(new_population) < pop_size:
            p1, p2 = random.sample(new_population, 2)
            c1, c2 = crossover(p1, p2)
            c1 = enforce_fixed_sensor_count(mutate(c1), sensor_num)
            c2 = enforce_fixed_sensor_count(mutate(c2), sensor_num)
            new_population.extend([c1, c2])

        population = np.array(new_population[:pop_size])

        if gen % 10 == 0 or gen == generations - 1:
            best_front = [fitness[i] for i in fronts[0]]
            print(f"Generation {gen}: Pareto front size = {len(best_front)}")
            for i, f in enumerate(best_front):
                print(f"  Sol {i+1}: VOI = {f[0]:.3f}, TE = {f[1]:.3f}")

    best_inds = [population[i] for i in fronts[0]]
    best_objs = [fitness[i] for i in fronts[0]]
    return best_inds, best_objs


def electre_selection(objs, weights=None):
    objs = np.array(objs)
    objs_norm = (objs - objs.min(axis=0)) / (objs.max(axis=0) - objs.min(axis=0) + 1e-9)
    objs_norm[:, 1] = 1 - objs_norm[:, 1]  # TE 为成本型，反转

    n = len(objs)
    weights = weights if weights is not None else np.array([0.5, 0.5])
    concordance = np.zeros((n, n))
    discordance = np.zeros((n, n))

    for i in range(n):
        for j in range(n):
            if i == j:
                continue
            better = objs_norm[i] >= objs_norm[j]
            worse = objs_norm[i] < objs_norm[j]
            concordance[i, j] = np.sum(weights[better])
            if np.any(worse):
                discordance[i, j] = np.max(np.abs(objs_norm[i] - objs_norm[j])[worse])
            else:
                discordance[i, j] = 0

    c_thresh = np.mean(concordance)
    d_thresh = np.mean(discordance)
    outrank = (concordance >= c_thresh) & (discordance <= d_thresh)
    scores = np.sum(outrank, axis=1)
    best_index = np.argmax(scores)
    return best_index, scores

voi_workbook= load_workbook(r"D:\A_科研相关\小论文2_合作监测点布置\覆盖双目标优化\L-town模拟验证\Khorshidi_selected_nodes_VOI_matrix.xlsx", data_only=True)
voi_worksheet=voi_workbook["Sheet1"]
voi_range =voi_worksheet["B2:CW101"]
voi_matrix= np.array([[cell.value for cell in row] for row in voi_range])

te_workbook= load_workbook(r"D:\A_科研相关\小论文2_合作监测点布置\覆盖双目标优化\L-town模拟验证\Khorshidi_selected_nodes_TE_matrix.xlsx", data_only=True)
te_worksheet=te_workbook["Sheet1"]
te_range =te_worksheet["B2:CW101"]
te_matrix= np.array([[cell.value for cell in row] for row in te_range])

clustered_nodes=(pd.read_excel(r"D:\A_科研相关\小论文2_合作监测点布置\覆盖双目标优化\L-town模拟验证\base_clustered_nodes.xlsx", sheet_name='Sheet1'))['index'].tolist()

# 你的 VOI_matrix 和 TE_matrix 是 N×N 的归一化矩阵
# 假设你已经有归一化后的 VOI_matrix 和 TE_matrix
sensor_num=8
for run in range(1,11):
    best_inds, best_objs = run_fixed_sensor_optimization(voi_matrix, te_matrix, sensor_num)

    for x, (voi, te) in zip(best_inds, best_objs):
        selected_nodes = np.where(x == 1)[0]
        print(f"Selected Nodes: {selected_nodes.tolist()}, VOI = {voi:.3f}, TE = {te:.3f}")
    # ELECTRE 多准则决策选最优
    best_idx, scores = electre_selection(best_objs)
    best_x = best_inds[best_idx]
    selected_nodes = np.where(best_x == 1)[0]
    actual_nodes = [clustered_nodes[i] for i in selected_nodes.tolist()]
    actual_nodes_df = pd.DataFrame([actual_nodes])  # 注意外面加 []
    file_path = os.path.join(r"D:\A_科研相关\小论文2_合作监测点布置\覆盖双目标优化\L-town模拟验证\算法不确定性", f"Khorshidi_{run}.xlsx")
    actual_nodes_df.to_excel(file_path , index=False, header=False)
# print(selected_nodes)
# print(f"\n✅ ELECTRE选中解: Nodes = {actual_nodes}, VOI = {best_objs[best_idx][0]:.3f}, TE = {best_objs[best_idx][1]:.3f}")

# VOIs = [voi for voi, te in best_objs]
# TEs = [te for voi, te in best_objs]

# plt.figure(figsize=(7, 5))
# plt.scatter(VOIs, TEs, c='blue', edgecolors='black', s=60)
# plt.xlabel("VOI信息价值", fontsize=12)
# plt.ylabel("TE信息冗余", fontsize=12)
# plt.title("Pareto 前沿 - VOI vs TE", fontsize=14)
# plt.grid(True)
# plt.gca().invert_xaxis()  # 因为 VOI 是越大越好，通常在左边
# plt.tight_layout()
# plt.show()

# list2=[737,262,672,734,561,74,100,521]#106,216,265,397,454,507,583,739 #391,463,50,52,432,635,443,397#616,443,641,683,516,336,422,352 #347,203,268,516,252,74,397,50
# individual=np.array([1 if item in list2 else 0 for item in clustered_nodes])
# voi, te=evaluate_individual(individual,voi_matrix, te_matrix)
# print(voi)
# print(te)
